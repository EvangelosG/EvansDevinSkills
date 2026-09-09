import copy

import pytest
from build_deck import DeckBuilder
from check_deck import Checker
from deckkit import TEMPLATE_PATH
from pptx import Presentation
from pptx.util import Inches, Pt

METRICS = "month,otd\n2026-01,91.2\n2026-02,90.4\n"


@pytest.fixture
def check(brand, tmp_path):
    def _check(slides, files=None, rules=None):
        for name, text in (files or {}).items():
            (tmp_path / name).write_text(text)
        spec_path = tmp_path / "deck.yaml"
        spec_path.touch()
        used = copy.deepcopy(brand)
        used["rules"].update(rules or {})
        prs = DeckBuilder(used, TEMPLATE_PATH).build({"slides": slides}, spec_path)
        return Checker(used, prs).run(), prs

    return _check


def test_the_worked_example_passes_every_rule(brand, example_deck):
    assert Checker(brand, Presentation(example_deck)).run() == []


def test_too_many_bullets_are_reported_with_the_limit(check, brand):
    limit = brand["rules"]["max_bullets_per_slide"]
    problems, _ = check(
        [{"type": "bullets", "title": "Summary", "bullets": [f"b{i}" for i in range(limit + 1)]}]
    )
    assert any(f"has {limit + 1} bullets, limit is {limit}" in p for p in problems)


def test_an_over_long_bullet_is_reported(check):
    problems, _ = check(
        [{"type": "bullets", "title": "Summary", "bullets": ["x" * 200, "short"]}],
        rules={"max_chars_per_bullet": 40},
    )
    assert any("bullet is 200 characters, limit is 40" in p for p in problems)


def test_an_over_long_title_is_reported(check):
    problems, _ = check(
        [{"type": "bullets", "title": "T" * 90, "bullets": ["one", "two"]}],
        rules={"max_title_chars": 70},
    )
    assert any("title is 90 characters, limit is 70" in p for p in problems)


def test_a_data_slide_without_a_source_note_is_reported(check):
    problems, _ = check(
        [{"type": "table", "title": "KPIs", "source": "m.csv"}], {"m.csv": METRICS}
    )
    assert any("no 'Source:' note" in p for p in problems)


def test_a_data_slide_with_a_source_note_is_accepted(check):
    problems, _ = check(
        [
            {
                "type": "table",
                "title": "KPIs",
                "source": "m.csv",
                "source_note": "Source: m.csv",
            }
        ],
        {"m.csv": METRICS},
    )
    assert problems == []


def test_a_missing_footer_is_reported(check, brand):
    problems, prs = check([{"type": "bullets", "title": "Summary", "bullets": ["one", "two"]}])
    assert problems == []

    slide = prs.slides[0]
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.text_frame.text == brand["footer"]:
            shape._element.getparent().remove(shape._element)
    assert any("no footer" in p for p in Checker(brand, prs).run())


def test_a_shape_hanging_off_the_slide_is_reported(check, brand):
    problems, prs = check([{"type": "bullets", "title": "Summary", "bullets": ["one", "two"]}])
    assert problems == []

    prs.slides[0].shapes.title.left = prs.slide_width - Inches(0.1)
    assert any("extends past the slide edge" in p for p in Checker(brand, prs).run())


def test_an_off_brand_font_is_reported(check, brand):
    problems, prs = check([{"type": "bullets", "title": "Summary", "bullets": ["one", "two"]}])
    assert problems == []

    for run in prs.slides[0].shapes.title.text_frame.paragraphs[0].runs:
        run.font.name = "Comic Sans MS"
    assert any("uses Comic Sans MS" in p for p in Checker(brand, prs).run())


def test_text_estimated_to_overflow_its_box_is_reported(check, brand):
    problems, prs = check([{"type": "bullets", "title": "Summary", "bullets": ["one", "two"]}])
    assert problems == []

    body = prs.slides[0].placeholders[1]
    body.height = Inches(0.3)
    for para in body.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(40)
    assert any("of height but the box is" in p for p in Checker(brand, prs).run())
