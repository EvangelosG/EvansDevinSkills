import pytest
from build_deck import DeckBuilder
from deckkit import TEMPLATE_PATH, DeckError
from pptx import Presentation


@pytest.fixture
def build(brand, tmp_path):
    def _build(slides, files=None):
        for name, text in (files or {}).items():
            (tmp_path / name).write_text(text)
        spec_path = tmp_path / "deck.yaml"
        spec_path.touch()
        return DeckBuilder(brand, TEMPLATE_PATH).build({"slides": slides}, spec_path)

    return _build


METRICS = "month,otd,cost\n2026-01,91.2,42.10\n2026-02,90.4,43.65\n"


def texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def test_every_slide_carries_the_footer_and_its_number(build, brand):
    prs = build(
        [
            {"type": "title", "title": "Q3 review"},
            {"type": "bullets", "title": "Summary", "bullets": ["one", "two"]},
        ]
    )
    for number, slide in enumerate(prs.slides, start=1):
        assert brand["footer"] in texts(slide)
        assert str(number) in texts(slide)


def test_title_slide_joins_subtitle_presenter_and_date(build):
    prs = build(
        [
            {
                "type": "title",
                "title": "Q3 review",
                "subtitle": "Northwind",
                "presenter": "Ops",
                "date": "October 2026",
            }
        ]
    )
    assert "Northwind\nOps\nOctober 2026" in texts(prs.slides[0])


def test_two_column_slide_puts_the_heading_above_its_bullets(build):
    prs = build(
        [
            {
                "type": "two_column",
                "title": "Drivers",
                "left_heading": "Helped",
                "left_bullets": ["automation"],
                "right_heading": "Hurt",
                "right_bullets": ["outage"],
            }
        ]
    )
    assert "Helped\nautomation" in texts(prs.slides[0])
    assert "Hurt\noutage" in texts(prs.slides[0])


def test_table_slide_reads_the_csv_and_keeps_the_header(build):
    prs = build(
        [{"type": "table", "title": "KPIs", "source": "m.csv"}], {"m.csv": METRICS}
    )
    table = next(s.table for s in prs.slides[0].shapes if s.has_table)
    assert [c.text for c in table.rows[0].cells] == ["month", "otd", "cost"]
    assert table.rows[1].cells[1].text == "91.2"


def test_table_slide_is_capped_at_max_table_rows(build, brand):
    rows = "\n".join(f"2026-{m:02d},9{m},4{m}" for m in range(1, 12))
    prs = build(
        [{"type": "table", "title": "KPIs", "source": "m.csv"}],
        {"m.csv": f"month,otd,cost\n{rows}\n"},
    )
    table = next(s.table for s in prs.slides[0].shapes if s.has_table)
    assert len(table.rows) == brand["rules"]["max_table_rows"] + 1


def test_chart_slide_builds_a_native_chart_from_the_named_columns(build):
    prs = build(
        [
            {
                "type": "chart",
                "title": "On-time delivery",
                "source": "m.csv",
                "category_column": "month",
                "series_columns": ["otd"],
                "chart_type": "line",
            }
        ],
        {"m.csv": METRICS},
    )
    chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert list(chart.plots[0].categories) == ["2026-01", "2026-02"]
    assert chart.plots[0].series[0].values == (91.2, 90.4)


def test_chart_slide_rejects_an_unsupported_chart_type(build):
    with pytest.raises(DeckError, match="chart_type 'donut' is not one of"):
        build(
            [
                {
                    "type": "chart",
                    "title": "On-time delivery",
                    "source": "m.csv",
                    "category_column": "month",
                    "series_columns": ["otd"],
                    "chart_type": "donut",
                }
            ],
            {"m.csv": METRICS},
        )


def test_a_slide_missing_a_required_key_names_the_slide_and_the_key(build):
    with pytest.raises(DeckError, match="slide 1 \\(chart\\) is missing 'category_column'"):
        build([{"type": "chart", "title": "Trend", "source": "m.csv"}], {"m.csv": METRICS})


def test_data_files_resolve_relative_to_the_spec_not_the_cwd(build, tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "m.csv").write_text(METRICS)
    prs = build([{"type": "table", "title": "KPIs", "source": "data/m.csv"}])
    assert any(s.has_table for s in prs.slides[0].shapes)


def test_unused_placeholders_are_dropped_so_they_do_not_render_as_prompts(build):
    prs = build([{"type": "section", "title": "Risks and asks"}])
    assert not any("Click to add" in t for t in texts(prs.slides[0]))


def test_the_deck_saves_and_reopens(build, tmp_path):
    prs = build([{"type": "bullets", "title": "Summary", "bullets": ["one"]}])
    out = tmp_path / "deck.pptx"
    prs.save(out)
    assert len(Presentation(out).slides) == 1
