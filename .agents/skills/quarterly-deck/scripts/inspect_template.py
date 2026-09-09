#!/usr/bin/env python3
"""List a template's layouts and placeholders, so LAYOUTS in deckkit.py can be re-mapped.

    python3 scripts/inspect_template.py [template.pptx]

Run this first after swapping in a corporate template: layout order is different
in every deck, and a wrong index is the usual reason a "branded" deck comes out
looking nothing like the brand.
"""

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from deckkit import LAYOUTS, TEMPLATE_PATH


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("template", nargs="?", type=Path, default=TEMPLATE_PATH)
    args = ap.parse_args()

    if not args.template.exists():
        raise SystemExit(f"error: {args.template} does not exist")

    prs = Presentation(args.template)
    print(
        f"{args.template.name}: {Emu(prs.slide_width).inches:.2f} x "
        f"{Emu(prs.slide_height).inches:.2f} in, {len(prs.slide_layouts)} layouts"
    )
    by_index = {index: name for name, index in LAYOUTS.items()}
    for i, layout in enumerate(prs.slide_layouts):
        mapped = f"  <- '{by_index[i]}'" if i in by_index else ""
        print(f"\n[{i}] {layout.name}{mapped}")
        for ph in layout.placeholders:
            fmt = ph.placeholder_format
            print(f"      idx={fmt.idx:<3} {str(fmt.type).split('.')[-1]:<16} {ph.name}")


if __name__ == "__main__":
    main()
