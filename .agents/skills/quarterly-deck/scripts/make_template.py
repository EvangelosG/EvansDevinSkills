#!/usr/bin/env python3
"""Generate a stand-in template.pptx from brand.yaml.

Only for teams that do not have a real corporate template yet. If yours exists,
drop it in as template.pptx instead and never run this — a real master carries
logos, picture placeholders and layout geometry this cannot invent.

    python3 scripts/make_template.py [--out template.pptx]
"""

import argparse
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from deckkit import SKILL_DIR, load_brand

# 16:9 at 13.333in x 7.5in, the PowerPoint default since 2013.
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

THEME_SLOTS = (
    ("dk1", "dark"),
    ("lt1", "light"),
    ("dk2", "dark"),
    ("lt2", "muted"),
    ("accent1", "accent1"),
    ("accent2", "accent2"),
    ("accent3", "accent3"),
    ("accent4", "accent4"),
    ("accent5", "muted"),
    ("accent6", "dark"),
    ("hlink", "accent1"),
    ("folHlink", "accent4"),
)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
THEME_RELTYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
)


def _theme_part(prs):
    """The slide master's theme part. python-pptx keeps it as an opaque blob."""
    return prs.slide_master.part.part_related_by(THEME_RELTYPE)


def _recolor(theme, colors):
    scheme = theme.find(f"{{{A}}}themeElements/{{{A}}}clrScheme")
    for slot, brand_key in THEME_SLOTS:
        node = scheme.find(f"{{{A}}}{slot}")
        if node is None:
            continue
        for child in list(node):
            node.remove(child)
        srgb = node.makeelement(f"{{{A}}}srgbClr", {"val": colors[brand_key].upper()})
        node.append(srgb)


def _refont(theme, fonts):
    scheme = theme.find(f"{{{A}}}themeElements/{{{A}}}fontScheme")
    for tag, key in (("majorFont", "heading"), ("minorFont", "body")):
        latin = scheme.find(f"{{{A}}}{tag}/{{{A}}}latin")
        if latin is not None:
            latin.set("typeface", fonts[key])


def _widen(prs, ratio: float):
    """python-pptx's default master is 4:3; widening the page leaves every
    placeholder in the left three quarters unless they are scaled too.

    All four values are written, not just the two that change: a layout
    placeholder normally inherits its geometry from the master, and setting
    only `left` leaves it with a half-written xfrm whose height is zero.
    """
    for master in prs.slide_masters:
        # Read every effective geometry before writing any of it: a layout
        # placeholder inherits from the master, so scaling the master first
        # would make the layouts read an already-scaled value and double it.
        geometry = []
        for container in [master] + list(master.slide_layouts):
            for shape in container.shapes:
                box = (shape.left, shape.top, shape.width, shape.height)
                if None not in box:
                    geometry.append((shape, box))

        for shape, (left, top, width, height) in geometry:
            shape.left = Emu(int(left * ratio))
            shape.width = Emu(int(width * ratio))
            shape.top = Emu(top)
            shape.height = Emu(height)


def _brand_the_master(prs, _brand: dict):
    """Titles start at the left margin, bold, at a fixed size.

    Set on the layouts rather than in build_deck.py on purpose: this is the
    part a real corporate .pptx would already carry.
    """
    for layout in prs.slide_master.slide_layouts:
        titles = [
            ph
            for ph in layout.placeholders
            if ph.placeholder_format.idx == 0 and ph.has_text_frame
        ]
        if not titles:
            continue
        paragraph = titles[0].text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.font.size = Pt(30)
        paragraph.font.bold = True


def build(out_path: Path, brand: dict) -> Path:
    prs = Presentation()
    ratio = SLIDE_W / prs.slide_width
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    _widen(prs, ratio)
    _brand_the_master(prs, brand)

    part = _theme_part(prs)
    theme = etree.fromstring(part.blob)
    _recolor(theme, brand["colors"])
    _refont(theme, brand["fonts"])
    part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)

    prs.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=SKILL_DIR / "template.pptx")
    args = ap.parse_args()

    brand = load_brand()
    path = build(args.out, brand)
    print(f"wrote {path} ({brand['company']} colors, {brand['fonts']['heading']})")


if __name__ == "__main__":
    main()
