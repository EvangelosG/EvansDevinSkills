#!/usr/bin/env python3
"""Render a deck spec (YAML) into a .pptx built on template.pptx.

    python3 scripts/build_deck.py example/deck.yaml --out /tmp/q3-review.pptx

Every slide is placed on a layout from the template, so the deck inherits the
master's theme instead of inventing one. Charts are native PowerPoint charts,
not images, so whoever receives the deck can still edit the numbers.
"""

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from deckkit import (
    BRAND_PATH,
    LAYOUTS,
    TEMPLATE_PATH,
    DeckError,
    load_brand,
    load_spec,
    read_csv,
    resolve,
    rgb,
    to_number,
)

CHART_TYPES = {
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
}

BODY_TOP = Inches(1.6)
BODY_LEFT = Inches(0.9)
FOOTER_HEIGHT = Inches(0.35)


def _fix_box(frame):
    """python-pptx creates textboxes with wrap off and autofit on, which lets a
    renderer resize the box around its own centre and undo the alignment."""
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE


class DeckBuilder:
    def __init__(self, brand: dict, template: Path):
        self.brand = brand
        self.prs = Presentation(template)
        self.width = self.prs.slide_width
        self.height = self.prs.slide_height

    # -- helpers ---------------------------------------------------------

    def _add(self, layout_name: str):
        return self.prs.slides.add_slide(self.prs.slide_layouts[LAYOUTS[layout_name]])

    def _strip_empty_placeholders(self, slide):
        """Unused placeholders render as 'Click to add text' prompts in exports."""
        for shape in list(slide.placeholders):
            if not shape.has_text_frame or shape.text_frame.text.strip():
                continue
            shape._element.getparent().remove(shape._element)

    def _set_title(self, slide, text: str, *, size=Pt(30), align=PP_ALIGN.LEFT):
        # Alignment and size are set on the run rather than left to the layout:
        # inherited paragraph properties are honoured inconsistently by
        # renderers other than PowerPoint, and the deck gets rendered.
        title = slide.shapes.title
        title.text = text
        for paragraph in title.text_frame.paragraphs:
            paragraph.alignment = align
            for run in paragraph.runs:
                run.font.size = size
                run.font.bold = True
                run.font.color.rgb = rgb(self.brand, "dark")
                run.font.name = self.brand["fonts"]["heading"]
        return title

    def _fill_bullets(self, frame, bullets, size=Pt(18)):
        frame.word_wrap = True
        for i, bullet in enumerate(bullets):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = str(bullet)
            para.level = 0
            para.alignment = PP_ALIGN.LEFT
            para.space_after = Pt(10)
            for run in para.runs:
                run.font.size = size
                run.font.name = self.brand["fonts"]["body"]
                run.font.color.rgb = rgb(self.brand, "dark")

    def _source_note(self, slide, text: str):
        box = slide.shapes.add_textbox(
            BODY_LEFT,
            self.height - Inches(1.0),
            self.width - BODY_LEFT * 2,
            Inches(0.3),
        )
        _fix_box(box.text_frame)
        para = box.text_frame.paragraphs[0]
        para.text = text
        para.alignment = PP_ALIGN.LEFT
        run = para.runs[0]
        run.font.size = Pt(10)
        run.font.italic = True
        run.font.name = self.brand["fonts"]["body"]
        run.font.color.rgb = rgb(self.brand, "muted")

    def _decorate(self, slide, number: int):
        """Accent bar, brand footer, slide number. Every slide, no exceptions."""
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), self.width, Inches(0.12)
        )
        bar.name = "Brand Bar"
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(self.brand, "accent1")
        bar.line.fill.background()
        bar.shadow.inherit = False

        top = self.height - FOOTER_HEIGHT - Inches(0.15)
        left_box = slide.shapes.add_textbox(
            BODY_LEFT, top, self.width / 2, FOOTER_HEIGHT
        )
        _fix_box(left_box.text_frame)
        para = left_box.text_frame.paragraphs[0]
        para.text = self.brand["footer"]
        para.alignment = PP_ALIGN.LEFT
        run = para.runs[0]
        run.font.size = Pt(9)
        run.font.name = self.brand["fonts"]["body"]
        run.font.color.rgb = rgb(self.brand, "muted")

        right_box = slide.shapes.add_textbox(
            self.width - Inches(1.6), top, Inches(0.9), FOOTER_HEIGHT
        )
        _fix_box(right_box.text_frame)
        para = right_box.text_frame.paragraphs[0]
        para.text = str(number)
        para.alignment = PP_ALIGN.RIGHT
        run = para.runs[0]
        run.font.size = Pt(9)
        run.font.name = self.brand["fonts"]["body"]
        run.font.color.rgb = rgb(self.brand, "muted")

    # -- slide types -----------------------------------------------------

    def title_slide(self, item, _spec_path):
        slide = self._add("title")
        self._set_title(slide, item["title"], size=Pt(40), align=PP_ALIGN.CENTER)
        subtitle_parts = [item.get("subtitle"), item.get("presenter"), item.get("date")]
        subtitle = "\n".join(p for p in subtitle_parts if p)
        if subtitle:
            slide.placeholders[1].text = subtitle
            for para in slide.placeholders[1].text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(16)
                    run.font.color.rgb = rgb(self.brand, "muted")
        return slide

    def section_slide(self, item, _spec_path):
        slide = self._add("section")
        title = self._set_title(slide, item["title"], size=Pt(34))
        for para in title.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = rgb(self.brand, "accent1")
        if item.get("text"):
            slide.placeholders[1].text = item["text"]
        return slide

    def bullets_slide(self, item, _spec_path):
        slide = self._add("bullets")
        self._set_title(slide, item["title"])
        self._fill_bullets(slide.placeholders[1].text_frame, item.get("bullets", []))
        return slide

    def two_column_slide(self, item, _spec_path):
        slide = self._add("two_column")
        self._set_title(slide, item["title"])
        for idx, side in ((1, "left"), (2, "right")):
            heading = item.get(f"{side}_heading")
            bullets = item.get(f"{side}_bullets", [])
            entries = ([heading] if heading else []) + list(bullets)
            frame = slide.placeholders[idx].text_frame
            self._fill_bullets(frame, entries, size=Pt(16))
            if heading:
                for run in frame.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = rgb(self.brand, "accent1")
        return slide

    def table_slide(self, item, spec_path):
        slide = self._add("title_only")
        self._set_title(slide, item["title"])

        header, rows = read_csv(
            resolve(spec_path, item["source"]), item.get("columns")
        )
        max_rows = self.brand["rules"]["max_table_rows"]
        rows = rows[:max_rows]

        width = self.width - BODY_LEFT * 2
        height = Inches(0.4) * (len(rows) + 1)
        shape = slide.shapes.add_table(
            len(rows) + 1, len(header), BODY_LEFT, BODY_TOP, width, height
        )
        table = shape.table

        for col, name in enumerate(header):
            cell = table.cell(0, col)
            cell.text = name
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(self.brand, "accent1")
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.bold = True
                    run.font.size = Pt(13)
                    run.font.name = self.brand["fonts"]["body"]
                    run.font.color.rgb = rgb(self.brand, "light")

        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(value)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(12)
                        run.font.name = self.brand["fonts"]["body"]
                        run.font.color.rgb = rgb(self.brand, "dark")

        if item.get("source_note"):
            self._source_note(slide, item["source_note"])
        return slide

    def chart_slide(self, item, spec_path):
        slide = self._add("title_only")
        self._set_title(slide, item["title"])

        category_col = item["category_column"]
        series_cols = item["series_columns"]
        data_path = resolve(spec_path, item["source"])
        header, rows = read_csv(data_path, [category_col] + list(series_cols))

        chart_data = CategoryChartData()
        chart_data.categories = [row[0] for row in rows]
        for i, name in enumerate(series_cols, start=1):
            chart_data.add_series(
                name,
                [to_number(row[i], where=f"{data_path.name}:{name}") for row in rows],
            )

        kind = item.get("chart_type", "line")
        if kind not in CHART_TYPES:
            raise DeckError(
                f"chart_type {kind!r} is not one of {', '.join(CHART_TYPES)}"
            )

        frame = slide.shapes.add_chart(
            CHART_TYPES[kind],
            BODY_LEFT,
            BODY_TOP,
            self.width - BODY_LEFT * 2,
            Inches(4.2),
            chart_data,
        )
        chart = frame.chart
        chart.has_title = False
        chart.font.size = Pt(12)
        chart.font.name = self.brand["fonts"]["body"]
        chart.font.color.rgb = rgb(self.brand, "dark")
        if len(series_cols) > 1:
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        else:
            chart.has_legend = False

        palette = ["accent1", "accent2", "accent3", "accent4"]
        for i, series in enumerate(chart.plots[0].series):
            color = rgb(self.brand, palette[i % len(palette)])
            if kind == "line":
                series.format.line.color.rgb = color
                series.format.line.width = Pt(2.5)
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color

        if item.get("source_note"):
            self._source_note(slide, item["source_note"])
        return slide

    def closing_slide(self, item, spec_path):
        slide = self.bullets_slide(item, spec_path)
        for para in slide.placeholders[1].text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = rgb(self.brand, "accent1")
        return slide

    # -- driver ----------------------------------------------------------

    def build(self, spec: dict, spec_path: Path) -> Presentation:
        handlers = {
            "title": self.title_slide,
            "section": self.section_slide,
            "bullets": self.bullets_slide,
            "two_column": self.two_column_slide,
            "table": self.table_slide,
            "chart": self.chart_slide,
            "closing": self.closing_slide,
        }
        for number, item in enumerate(spec["slides"], start=1):
            try:
                slide = handlers[item["type"]](item, spec_path)
            except KeyError as exc:
                raise DeckError(
                    f"slide {number} ({item['type']}) is missing {exc}"
                ) from exc
            self._strip_empty_placeholders(slide)
            self._decorate(slide, number)
        return self.prs


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("spec", type=Path, help="deck spec YAML")
    ap.add_argument("--out", type=Path, required=True, help="output .pptx")
    ap.add_argument("--template", type=Path, default=TEMPLATE_PATH)
    ap.add_argument("--brand", type=Path, default=BRAND_PATH, help="brand.yaml to use")
    args = ap.parse_args()

    if not args.template.exists():
        raise SystemExit(
            f"{args.template} does not exist — drop in your corporate template, or run "
            "python3 scripts/make_template.py to generate a stand-in"
        )

    brand = load_brand(args.brand)
    spec = load_spec(args.spec)
    prs = DeckBuilder(brand, args.template).build(spec, args.spec.resolve())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    print(f"wrote {args.out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    try:
        main()
    except DeckError as exc:
        raise SystemExit(f"error: {exc}")
