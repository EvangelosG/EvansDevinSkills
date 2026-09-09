#!/usr/bin/env python3
"""Check a built .pptx against brand.yaml. Exits non-zero on any violation.

    python3 scripts/check_deck.py /tmp/q3-review.pptx

What it catches is mechanical: too many bullets, over-long titles, a missing
footer or source note, shapes hanging off the slide, fonts that are not the
brand's, and text that is *estimated* to overflow its box.

What it cannot catch is whether the deck says anything worth saying, or whether
a chart is misleading. Render the slides and look at them as well — the
estimate below is a proxy for a rasteriser, not a substitute for eyes.
"""

import argparse
import math
import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu

from deckkit import BRAND_PATH, DeckError, load_brand

# Any TrueType face gives a better width estimate than counting characters;
# this one ships with Pillow's usual dependencies on Linux boxes.
FALLBACK_FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
LINE_SPACING = 1.2
DEFAULT_BODY_PT = 18.0


def _font(size_pt: float):
    try:
        return ImageFont.truetype(FALLBACK_FONT, int(round(size_pt)))
    except OSError:
        return None


def _text_width_pt(text: str, size_pt: float) -> float:
    font = _font(size_pt)
    if font is None:
        return len(text) * size_pt * 0.5
    return font.getlength(text)


def _para_size_pt(paragraph) -> float:
    for run in paragraph.runs:
        if run.font.size is not None:
            return run.font.size.pt
    if paragraph.font.size is not None:
        return paragraph.font.size.pt
    return DEFAULT_BODY_PT


def estimate_overflow(shape) -> tuple[float, float] | None:
    """Return (needed_pt, available_pt) when text is estimated to overflow."""
    frame = shape.text_frame
    if not frame.text.strip() or shape.width is None or shape.height is None:
        return None

    inset = (frame.margin_left or 0) + (frame.margin_right or 0)
    usable_pt = max(Emu(shape.width - inset).pt, 1)
    needed_pt = 0.0
    for paragraph in frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs)
        size_pt = _para_size_pt(paragraph)
        if not text:
            needed_pt += size_pt * LINE_SPACING
            continue
        lines = max(1, math.ceil(_text_width_pt(text, size_pt) / usable_pt))
        needed_pt += lines * size_pt * LINE_SPACING
        if paragraph.space_after is not None:
            needed_pt += paragraph.space_after.pt

    vertical_inset = (frame.margin_top or 0) + (frame.margin_bottom or 0)
    available_pt = Emu(shape.height - vertical_inset).pt
    if needed_pt > available_pt * 1.02:  # 2% slack: this is an estimate
        return needed_pt, available_pt
    return None


class Checker:
    def __init__(self, brand: dict, prs: Presentation):
        self.brand = brand
        self.rules = brand["rules"]
        self.prs = prs
        self.problems: list[str] = []

    def fail(self, slide_no: int, message: str):
        self.problems.append(f"slide {slide_no}: {message}")

    def run(self) -> list[str]:
        allowed_fonts = {self.brand["fonts"]["heading"], self.brand["fonts"]["body"]}
        for number, slide in enumerate(self.prs.slides, start=1):
            self._check_footer(number, slide)
            self._check_data_note(number, slide)
            for shape in slide.shapes:
                self._check_bounds(number, slide, shape)
                if not shape.has_text_frame:
                    continue
                self._check_fonts(number, shape, allowed_fonts)
                self._check_overflow(number, shape)
            self._check_title(number, slide)
            self._check_bullets(number, slide)
        return self.problems

    def _texts(self, slide):
        return [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]

    def _check_footer(self, number, slide):
        if not self.rules.get("require_footer"):
            return
        if not any(self.brand["footer"] in text for text in self._texts(slide)):
            self.fail(number, f"no footer reading {self.brand['footer']!r}")

    def _check_data_note(self, number, slide):
        if not self.rules.get("require_source_note_on_data_slides"):
            return
        has_data = any(
            shape.has_table or shape.has_chart for shape in slide.shapes
        )
        if has_data and not any("source" in t.lower() for t in self._texts(slide)):
            self.fail(number, "has a table or chart but no 'Source:' note")

    def _check_bounds(self, number, slide, shape):
        if shape.left is None or shape.top is None:
            return
        right, bottom = shape.left + (shape.width or 0), shape.top + (shape.height or 0)
        if shape.left < 0 or shape.top < 0:
            self.fail(number, f"{shape.shape_type} '{shape.name}' starts off-slide")
        elif right > self.prs.slide_width or bottom > self.prs.slide_height:
            self.fail(number, f"'{shape.name}' extends past the slide edge")

    def _check_fonts(self, number, shape, allowed):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                name = run.font.name
                if name is not None and name not in allowed:
                    self.fail(
                        number,
                        f"'{shape.name}' uses {name}, not {' or '.join(sorted(allowed))}",
                    )
                    return

    def _check_overflow(self, number, shape):
        result = estimate_overflow(shape)
        if result is None:
            return
        needed, available = result
        self.fail(
            number,
            f"'{shape.name}' text needs ~{needed:.0f}pt of height but the box is "
            f"{available:.0f}pt — shorten it or split the slide",
        )

    def _check_title(self, number, slide):
        title = slide.shapes.title
        if title is None or not title.text_frame.text.strip():
            return
        limit = self.rules["max_title_chars"]
        length = len(title.text_frame.text.strip())
        if length > limit:
            self.fail(number, f"title is {length} characters, limit is {limit}")

    def _check_bullets(self, number, slide):
        max_bullets = self.rules["max_bullets_per_slide"]
        max_chars = self.rules["max_chars_per_bullet"]
        title = slide.shapes.title
        for shape in slide.shapes:
            if not shape.has_text_frame or shape is title:
                continue
            paragraphs = [
                p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()
            ]
            if len(paragraphs) < 2:  # footers, source notes, single captions
                continue
            if len(paragraphs) > max_bullets:
                self.fail(
                    number,
                    f"'{shape.name}' has {len(paragraphs)} bullets, limit is {max_bullets}",
                )
            for paragraph in paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if len(text) > max_chars:
                    self.fail(
                        number,
                        f"bullet is {len(text)} characters, limit is {max_chars}: "
                        f"{text[:60]}...",
                    )


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("deck", type=Path)
    ap.add_argument("--brand", type=Path, default=BRAND_PATH, help="brand.yaml to use")
    args = ap.parse_args()

    if not args.deck.exists():
        raise SystemExit(f"error: {args.deck} does not exist")

    brand = load_brand(args.brand)
    problems = Checker(brand, Presentation(args.deck)).run()
    if problems:
        print(f"{len(problems)} problem(s) in {args.deck.name}:", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        raise SystemExit(1)
    print(f"{args.deck.name}: passes every rule in brand.yaml")


if __name__ == "__main__":
    try:
        main()
    except DeckError as exc:
        raise SystemExit(f"error: {exc}")
